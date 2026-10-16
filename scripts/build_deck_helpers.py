import os
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_DIR = Path(r"c:\Users\shrey\OneDrive\Desktop\We_Are_The_Bugs_Pitch_Deliverables")
IMAGES_DIR = OUTPUT_DIR / "Selected_Images"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

PPTX_PATH = OUTPUT_DIR / "We_Are_The_Bugs_Hackathon_Pitch.pptx"

# Initialize Presentation in 16:9 Widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

# Color Palette Constants
BG_COLOR = RGBColor(7, 13, 10)         # Deep Forest Black #070d0a
SURFACE_COLOR = RGBColor(14, 24, 19)   # Elevated Card #0e1813
SURFACE_BORDER = RGBColor(25, 45, 34)  # Subtle Green Border
ACCENT_GREEN = RGBColor(0, 229, 153)   # Emerald Green #00e599
ACCENT_CYAN = RGBColor(56, 189, 248)   # Tech Cyan #38bdf8
ACCENT_ORANGE = RGBColor(249, 115, 22) # Alert Orange #f97316
TEXT_WHITE = RGBColor(248, 250, 252)   # Pure Off-White
TEXT_MUTED = RGBColor(148, 163, 184)   # Slate Muted Gray
TEXT_DIM = RGBColor(100, 116, 139)     # Dim Gray

def apply_background(slide):
    """Draws full slide dark background."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_COLOR
    bg.line.fill.background()
    return bg

def add_header(slide, tracker: str, title: str, subtitle: str):
    """Adds uniform slide header with category badge, title, and subtitle."""
    # Top Tracker pill
    tb_tr = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
    tf_tr = tb_tr.text_frame
    tf_tr.word_wrap = True
    tf_tr.margin_left = tf_tr.margin_top = tf_tr.margin_right = tf_tr.margin_bottom = 0
    p_tr = tf_tr.paragraphs[0]
    p_tr.text = tracker.upper()
    p_tr.font.name = "Arial"
    p_tr.font.size = Pt(10)
    p_tr.font.bold = True
    p_tr.font.color.rgb = ACCENT_GREEN

    # Title
    tb_ti = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.55))
    tf_ti = tb_ti.text_frame
    tf_ti.word_wrap = True
    tf_ti.margin_left = tf_ti.margin_top = tf_ti.margin_right = tf_ti.margin_bottom = 0
    p_ti = tf_ti.paragraphs[0]
    p_ti.text = title
    p_ti.font.name = "Arial"
    p_ti.font.size = Pt(22)
    p_ti.font.bold = True
    p_ti.font.color.rgb = TEXT_WHITE

    # Subtitle
    if subtitle:
        tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.733), Inches(0.35))
        tf_sub = tb_sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(12)
        p_sub.font.color.rgb = TEXT_MUTED

def add_footer(slide, current_slide: int, total_slides: int = 13):
    """Uniform footer on all slides."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"ECOSYNC // SMART WASTE MANAGEMENT & DIGITAL TWIN          |          TEAM: WE ARE THE BUGS          |          SLIDE {current_slide:02d} OF {total_slides:02d}"
    p.font.name = "Arial"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = TEXT_DIM

def add_card(slide, left, top, width, height, fill_color=SURFACE_COLOR, border_color=SURFACE_BORDER):
    """Creates a rounded container card for visual grouping."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def add_speaker_notes(slide, script: str, main_point: str, transition: str, question: str, answer: str):
    """Embeds structured speaker notes into slide."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    notes_content = f"""🎙️ PRESENTER SCRIPT:
{script}

🎯 MAIN POINT TO EMPHASIZE:
{main_point}

🔄 TRANSITION TO NEXT SLIDE:
{transition}

❓ LIKELY JUDGE QUESTION:
"{question}"

💡 EVIDENCE-BASED ANSWER:
{answer}
"""
    text_frame.text = notes_content

print("Helper functions initialized.")
